"""Text extraction from PDF, DOCX and plain-text resumes."""
from __future__ import annotations

import io
import re
from pathlib import Path
from typing import BinaryIO, Union

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class UnsupportedFileError(ValueError):
    """Raised when a file type cannot be parsed."""


def _pdf_text(data: bytes) -> str:
    import pdfplumber

    pages = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text() or "")
    return "\n".join(pages)


def _docx_text(data: bytes) -> str:
    import docx  # python-docx

    document = docx.Document(io.BytesIO(data))
    lines = [p.text for p in document.paragraphs]
    # Tables are common in resumes (skills grids, education rows)
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def clean_text(text: str) -> str:
    """Normalise whitespace and common PDF artefacts while keeping line breaks."""
    text = text.replace("\u2022", "-").replace("\u25cf", "-").replace("\uf0b7", "-")
    text = text.replace("\u2013", "-").replace("\u2014", "-").replace("\xa0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()


def extract_text(source: Union[str, Path, bytes, BinaryIO], filename: str | None = None) -> str:
    """Extract clean text from a resume.

    ``source`` can be a path, raw bytes or a file-like object. When bytes or a
    stream are passed, ``filename`` is used to detect the format.
    """
    if isinstance(source, (str, Path)):
        path = Path(source)
        filename = filename or path.name
        data = path.read_bytes()
    elif isinstance(source, bytes):
        data = source
    else:
        data = source.read()

    if not filename:
        raise UnsupportedFileError("filename is required to detect the file type")
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        text = _pdf_text(data)
    elif ext == ".docx":
        text = _docx_text(data)
    elif ext == ".txt":
        text = data.decode("utf-8", errors="ignore")
    else:
        raise UnsupportedFileError(f"Unsupported file type '{ext}'. Use PDF, DOCX or TXT.")
    return clean_text(text)
