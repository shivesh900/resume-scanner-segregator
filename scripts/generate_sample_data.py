"""Generate the synthetic sample resumes in sample_data/resumes.

All people, emails, phone numbers and employers here are FICTIONAL. Each resume
is written to look like a real student/early-career CV so the parser has to deal
with sections, bullets and date ranges. ``expected_priority`` is the bucket a
human recruiter would put the candidate in for the bundled ML Engineer JD; it is
used by scripts/evaluate.py.

Run:  python scripts/generate_sample_data.py
"""
from __future__ import annotations

import csv
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "sample_data" / "resumes"

CANDIDATES = [
    dict(file="aarav_mehta.pdf", expected="High", name="Aarav Mehta", email="aarav.mehta@example.com",
         phone="+91 90000 11101", links="github.com/aarav-ml | linkedin.com/in/aarav-mehta-demo",
         summary="Machine learning engineer focused on NLP and production ML systems.",
         skills="Python, SQL, PyTorch, scikit-learn, Pandas, NumPy, NLP, Hugging Face Transformers, FastAPI, REST APIs, Docker, AWS, Git, MLflow",
         experience=[("ML Engineer Intern - Nimbus Analytics (fictional)", "Jun 2023 - Present",
                      ["Fine-tuned BERT models in PyTorch for ticket classification (F1 0.91)",
                       "Served models behind FastAPI REST endpoints in Docker on AWS"]),
                     ("Data Science Intern - Orbit Labs (fictional)", "Jan 2022 - May 2023",
                      ["Built scikit-learn churn models with Pandas/NumPy feature pipelines",
                       "Wrote SQL ETL jobs and tracked experiments with MLflow"])],
         education="B.Tech Computer Science, Example Institute of Technology, 2024 - CGPA 8.9",
         projects=["Resume NER with spaCy and PyTorch", "RAG chatbot over course notes"]),
    dict(file="priya_raman.docx", expected="High", name="Priya Raman", email="priya.raman@example.org",
         phone="+91 90000 11102", links="github.com/priya-r-demo",
         summary="AI/ML graduate with 2 years of internship experience in NLP and deep learning.",
         skills="Python, PyTorch, TensorFlow, scikit-learn, NLP, Machine Learning, Pandas, NumPy, SQL, PostgreSQL, Flask, REST API, Git, Linux",
         experience=[("NLP Research Intern - Lumen AI (fictional)", "Jul 2022 - Aug 2024",
                      ["Trained transformer sentiment models in PyTorch on 200k reviews",
                       "Exposed inference as a Flask REST API used by the product team"])],
         education="M.Tech Artificial Intelligence, Sample University, 2024",
         projects=["Multilingual sentiment classifier", "Topic modelling dashboard"]),
    dict(file="daniel_okafor.pdf", expected="High", name="Daniel Okafor", email="d.okafor@example.net",
         phone="+234 800 000 1103", links="linkedin.com/in/daniel-okafor-demo",
         summary="Machine Learning Engineer with 3 years of experience building NLP pipelines.",
         skills="Python, SQL, scikit-learn, PyTorch, NLP, spaCy, Pandas, NumPy, Git, REST APIs, Docker, Kubernetes, GCP",
         experience=[("Machine Learning Engineer - Delta Retail Tech (fictional)", "Mar 2021 - Apr 2024",
                      ["Shipped product-categorisation models (scikit-learn, PyTorch) to 2M SKUs",
                       "Maintained REST APIs and CI/CD for model releases with Docker and Kubernetes"])],
         education="B.Sc Computer Science, Demo State University, 2020",
         projects=["Search relevance ranking with learning-to-rank"]),
    dict(file="meera_krishnan.pdf", expected="High", name="Meera Krishnan", email="meera.k@example.com",
         phone="+91 90000 11104", links="github.com/meera-k-demo",
         summary="Final-year CSE student, ML intern, Kaggle expert in NLP competitions.",
         skills="Python, Machine Learning, Deep Learning, NLP, PyTorch, scikit-learn, Pandas, NumPy, SQL, MySQL, Git, Flask, REST APIs, AWS",
         experience=[("Machine Learning Intern - CloudNine Health (fictional)", "May 2023 - Jul 2024",
                      ["Built clinical-note classifiers with PyTorch and scikit-learn",
                       "Deployed a Flask REST API on AWS EC2"])],
         education="B.E. Computer Science and Engineering, Example College of Engineering, 2025",
         projects=["Fake news detection with BERT", "Movie recommender with collaborative filtering"]),
    dict(file="rohan_gupta.pdf", expected="Medium", name="Rohan Gupta", email="rohan.gupta@example.com",
         phone="+91 90000 11105", links="linkedin.com/in/rohan-gupta-demo",
         summary="Data analyst who turns messy data into dashboards and decisions.",
         skills="Python, SQL, MySQL, Pandas, NumPy, Power BI, Tableau, Statistics, Excel, Machine Learning basics",
         experience=[("Data Analyst - BrightCart (fictional)", "Aug 2022 - Present",
                      ["Built Power BI dashboards on SQL data marts for 40 stakeholders",
                       "Ran A/B test analysis with Pandas and statistics"])],
         education="B.Com with Data Analytics, Sample College, 2022",
         projects=["Sales forecasting with linear regression"]),
    dict(file="sara_thomas.docx", expected="Medium", name="Sara Thomas", email="sara.thomas@example.org",
         phone="+91 90000 11106", links="github.com/sara-t-demo",
         summary="Backend developer interested in moving into ML engineering.",
         skills="Python, Flask, Django, REST APIs, SQL, PostgreSQL, Redis, Docker, Git, Linux, Pandas",
         experience=[("Backend Developer - Paystream (fictional)", "Jan 2022 - Present",
                      ["Designed REST APIs in Flask and Django serving 1M requests/day",
                       "Containerised services with Docker"])],
         education="B.Tech Information Technology, Example Institute of Technology, 2021",
         projects=["Completed Andrew Ng's Machine Learning course; built a spam classifier with scikit-learn"]),
    dict(file="karthik_iyer.pdf", expected="Medium", name="Karthik Iyer", email="karthik.iyer@example.com",
         phone="+91 90000 11107", links="github.com/karthik-cv-demo",
         summary="Computer vision enthusiast working with deep learning models.",
         skills="Python, PyTorch, OpenCV, Deep Learning, NumPy, Matplotlib, Git",
         experience=[("Computer Vision Intern - Visionary Robotics (fictional)", "Jan 2024 - Jun 2024",
                      ["Trained YOLO object detectors in PyTorch for warehouse robots",
                       "Wrote OpenCV pre-processing for camera feeds"])],
         education="B.Tech Electronics and Communication, Demo University, 2025",
         projects=["Traffic sign recognition with CNNs"]),
    dict(file="ananya_das.docx", expected="Medium", name="Ananya Das", email="ananya.das@example.com",
         phone="+91 90000 11108", links="github.com/ananya-d-demo",
         summary="Fresh graduate passionate about machine learning and data science.",
         skills="Python, Machine Learning, scikit-learn, Pandas, NumPy, NLP, SQL, Git, Statistics",
         experience=[],
         education="B.Sc Statistics, Sample University, 2025",
         projects=["Twitter sentiment analysis with TF-IDF and logistic regression",
                   "House price prediction with scikit-learn"]),
    dict(file="vikram_singh.pdf", expected="Low", name="Vikram Singh", email="vikram.singh@example.net",
         phone="+91 90000 11109", links="github.com/vikram-fe-demo",
         summary="Frontend developer who loves building fast, accessible web apps.",
         skills="JavaScript, TypeScript, React, Next.js, HTML, CSS, Tailwind, Node.js, Figma, Git",
         experience=[("Frontend Developer - PixelForge (fictional)", "Jun 2021 - Present",
                      ["Built React and Next.js storefronts with Tailwind CSS",
                       "Improved Lighthouse performance score from 62 to 95"])],
         education="B.E. Computer Science, Example College of Engineering, 2021",
         projects=["Open-source React component library"]),
    dict(file="fatima_sheikh.pdf", expected="Low", name="Fatima Sheikh", email="fatima.sheikh@example.com",
         phone="+91 90000 11110", links="linkedin.com/in/fatima-sheikh-demo",
         summary="UI/UX designer crafting clean, user-centred mobile experiences.",
         skills="Figma, Adobe XD, UI/UX, Wireframing, Prototyping, User Research, HTML, CSS, Communication",
         experience=[("UI/UX Designer - Studio Nine (fictional)", "Apr 2022 - Present",
                      ["Led wireframing and prototyping for 3 fintech apps in Figma",
                       "Ran usability tests with 50+ users"])],
         education="B.Des Interaction Design, Demo Institute of Design, 2022",
         projects=["Redesign of a campus food-ordering app"]),
    dict(file="arjun_nair.docx", expected="Low", name="Arjun Nair", email="arjun.nair@example.org",
         phone="+91 90000 11111", links="linkedin.com/in/arjun-nair-demo",
         summary="Network and systems administrator with strong Linux skills.",
         skills="Linux, Bash, Networking, Cisco, Java, MySQL, Windows Server",
         experience=[("Systems Administrator - Metro Telecom (fictional)", "Jul 2019 - Present",
                      ["Managed 300 Linux servers and Cisco network equipment",
                       "Automated backups with Bash scripts"])],
         education="Diploma in Computer Engineering, Sample Polytechnic, 2019",
         projects=["Internal ticketing tool in Java"]),
    dict(file="lakshmi_pillai.pdf", expected="Low", name="Lakshmi Pillai", email="lakshmi.p@example.com",
         phone="+91 90000 11112", links="linkedin.com/in/lakshmi-pillai-demo",
         summary="Marketing executive with a flair for storytelling and campaigns.",
         skills="Content Marketing, SEO, Social Media, Communication, Leadership, Excel",
         experience=[("Marketing Executive - BlueLeaf Foods (fictional)", "May 2020 - Present",
                      ["Grew Instagram following from 5k to 60k",
                       "Managed a campaign budget of INR 20 lakh"])],
         education="MBA Marketing, Example Business School, 2020",
         projects=["Brand relaunch campaign"]),
]


def _lines(c):
    yield ("title", c["name"].upper())
    yield ("text", f"{c['email']} | {c['phone']} | {c['links']}")
    yield ("head", "Summary")
    yield ("text", c["summary"])
    yield ("head", "Skills")
    yield ("text", c["skills"])
    if c["experience"]:
        yield ("head", "Experience")
        for role, dates, bullets in c["experience"]:
            yield ("bold", f"{role}    {dates}")
            for b in bullets:
                yield ("bullet", b)
    yield ("head", "Education")
    yield ("text", c["education"])
    yield ("head", "Projects")
    for p in c["projects"]:
        yield ("bullet", p)


def write_pdf(c, path):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    ss = getSampleStyleSheet()
    flow = []
    for kind, text in _lines(c):
        if kind == "title":
            flow.append(Paragraph(text, ss["Title"]))
        elif kind == "head":
            flow += [Spacer(1, 6), Paragraph(text, ss["Heading2"])]
        elif kind == "bold":
            flow.append(Paragraph(f"<b>{text}</b>", ss["Normal"]))
        elif kind == "bullet":
            flow.append(Paragraph(text, ss["Normal"], bulletText="\u2022"))
        else:
            flow.append(Paragraph(text, ss["Normal"]))
    SimpleDocTemplate(str(path), pagesize=A4, title=f"{c['name']} - Resume",
                      author="synthetic sample", invariant=1).build(flow)


def write_docx(c, path):
    import docx

    d = docx.Document()
    for kind, text in _lines(c):
        if kind == "title":
            d.add_heading(text, level=0)
        elif kind == "head":
            d.add_heading(text, level=1)
        elif kind == "bold":
            d.add_paragraph().add_run(text).bold = True
        elif kind == "bullet":
            d.add_paragraph(text, style="List Bullet")
        else:
            d.add_paragraph(text)
    d.core_properties.author = "synthetic sample"
    d.save(str(path))


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    for c in CANDIDATES:
        path = OUT / c["file"]
        (write_pdf if path.suffix == ".pdf" else write_docx)(c, path)
    with open(ROOT / "sample_data" / "labels.csv", "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["filename", "expected_priority"])
        for c in CANDIDATES:
            w.writerow([c["file"], c["expected"]])
    print(f"Wrote {len(CANDIDATES)} synthetic resumes to {OUT}")


if __name__ == "__main__":
    main()
